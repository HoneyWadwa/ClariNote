"""Shared fixtures for the backend test-suite.

The OpenAI API is never called: tests monkeypatch ``generation.call_openai``.
"""

from __future__ import annotations

import io
import sys
from pathlib import Path
from typing import Any, Callable, Dict

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:  # allow `pytest` from any working directory
    sys.path.insert(0, str(ROOT))

from backend.app import create_app  # noqa: E402
from backend.config import Config  # noqa: E402

SAMPLE_TEXT = (ROOT / "sample" / "sample_input.txt").read_text(encoding="utf-8")


class FakeRequest:
    """Minimal stand-in for the request object attached to OpenAI exceptions."""

    method = "POST"
    url = "https://api.openai.com/v1/chat/completions"


class FakeResponse:
    """Minimal stand-in for the response object attached to OpenAI exceptions."""

    def __init__(self, status_code: int):
        self.status_code = status_code
        self.headers: dict = {}
        self.request = FakeRequest()


def fake_response(status_code: int) -> FakeResponse:
    return FakeResponse(status_code)

S1 = "Process scheduling is an operating system function that selects a process from the ready queue for execution by the CPU."
S2 = "First-Come, First-Served (FCFS) schedules processes in the order they arrive."
S3 = "Shortest Job First (SJF) selects the process with the shortest CPU burst."
S4 = "Round Robin assigns each process a fixed time quantum and cycles through the ready queue."
S5 = "A context switch occurs when the CPU changes from executing one process or thread to another."
S6 = "Scheduling algorithms can differ in waiting time, turnaround time, and response time."


def evidence(*quotes: str, location: str | None = None) -> list[Dict[str, Any]]:
    return [{"quote": quote, "location": location} for quote in quotes]


def valid_materials() -> Dict[str, Any]:
    """A schema-valid generation result with EXACTLY five questions."""
    return {
        "summary": [
            {"point": "Process scheduling selects a process from the ready queue for CPU execution.",
             "evidence": evidence(S1), "grounding": "SUPPORTED", "warning": None},
            {"point": "Round Robin provides the fastest response time of all algorithms.",
             "evidence": evidence(S6), "grounding": "UNSUPPORTED",
             "warning": "The notes do not rank algorithms."},
        ],
        "glossary": [
            {"term": "Process scheduling", "definition": "Selects a process from the ready queue for execution by the CPU.",
             "evidence": evidence(S1), "grounding": "SUPPORTED", "warning": None},
            {"term": "Turnaround time", "definition": None, "evidence": evidence(S6),
             "grounding": "UNSUPPORTED", "warning": "Not supported by source."},
        ],
        "questions": [
            {"number": 1, "question_type": "recall",
             "question": "What selects a process from the ready queue for execution by the CPU?",
             "evidence": evidence(S1), "grounding": "SUPPORTED",
             "answer": "Process scheduling.",
             "answer_evidence": evidence(S1), "answer_grounding": "SUPPORTED", "warning": None},
            {"number": 2, "question_type": "recall",
             "question": "How does FCFS order the processes it runs?",
             "evidence": evidence(S2), "grounding": "SUPPORTED",
             "answer": "It schedules processes in the order they arrive.",
             "answer_evidence": evidence(S2), "answer_grounding": "SUPPORTED", "warning": None},
            {"number": 3, "question_type": "conceptual",
             "question": "Which process does Shortest Job First select?",
             "evidence": evidence(S3), "grounding": "SUPPORTED",
             "answer": "The process with the shortest CPU burst.",
             "answer_evidence": evidence(S3), "answer_grounding": "SUPPORTED", "warning": None},
            {"number": 4, "question_type": "conceptual",
             "question": "What is a context switch?",
             "evidence": evidence(S5), "grounding": "SUPPORTED",
             "answer": "The CPU changes from executing one process or thread to another.",
             "answer_evidence": evidence(S5), "answer_grounding": "SUPPORTED", "warning": None},
            {"number": 5, "question_type": "application",
             "question": "Which three measures can scheduling algorithms differ in?",
             "evidence": evidence(S6), "grounding": "SUPPORTED",
             "answer": "Waiting time, turnaround time, and response time.",
             "answer_evidence": evidence(S6), "answer_grounding": "SUPPORTED", "warning": None},
        ],
    }


def build_pdf(text: str, lines_per_page: int = 8) -> bytes:
    """Create an in-memory PDF containing the given text, split over pages."""
    import pymupdf

    words, lines, current = text.split(), [], ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if len(candidate) > 68:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)

    document = pymupdf.open()
    page = document.new_page()
    y = 72
    for index, line in enumerate(lines):
        if index and index % lines_per_page == 0:
            page = document.new_page()
            y = 72
        page.insert_text((72, y), line)
        y += 15
    data = document.tobytes()
    document.close()
    return data


def build_docx(text: str) -> bytes:
    """Create an in-memory .docx with a heading followed by the given paragraphs."""
    import docx

    document = docx.Document()
    document.add_heading("Operating Systems Lecture", level=1)
    for paragraph in text.split("\n"):
        if paragraph.strip():
            document.add_paragraph(paragraph.strip())
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def sample_text() -> str:
    return SAMPLE_TEXT


@pytest.fixture
def make_app() -> Callable[..., Any]:
    def _make(**overrides: Any):
        config = Config(
            openai_api_key=overrides.get("api_key"),
            openai_model=overrides.get("model", "gpt-4o-mini"),
            max_file_mb=overrides.get("max_file_mb", 5),
            max_text_length=overrides.get("max_text_length", 60_000),
        )
        app = create_app(config)
        app.config["TESTING"] = True
        return app

    return _make


@pytest.fixture
def client(make_app):
    """Client with NO API key -> demo mode."""
    return make_app().test_client()


@pytest.fixture
def live_client(make_app):
    """Client configured as if an API key were present (AI calls are mocked)."""
    return make_app(api_key="sk-test-not-a-real-key").test_client()


@pytest.fixture
def materials() -> Dict[str, Any]:
    return valid_materials()


@pytest.fixture
def mock_openai(monkeypatch) -> Callable[[Any], None]:
    """Replace the OpenAI transport with a function returning fixed JSON."""

    def _install(payload: Any) -> None:
        from backend.services import generation

        def _call(*args, **kwargs):
            if isinstance(payload, Exception):
                raise payload
            return payload

        monkeypatch.setattr(generation, "call_openai", _call)

    return _install
